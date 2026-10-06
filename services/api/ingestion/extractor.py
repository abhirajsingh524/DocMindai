import os
from pathlib import Path
from typing import List, Dict, Any, Tuple
import pypdf
import docx

class ExtractedSection:
    def __init__(self, text: str, page_number: int | None = None, section_title: str | None = None):
        self.text = text
        self.page_number = page_number
        self.section_title = section_title

def extract_from_file(file_path: str, file_type: str) -> Tuple[List[ExtractedSection], List[str]]:
    warnings = []
    sections: List[ExtractedSection] = []
    
    normalized_type = file_type.lower().strip().replace(".", "")
    path = Path(file_path)
    
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")
        
    if normalized_type == "pdf":
        try:
            reader = pypdf.PdfReader(file_path)
            if reader.is_encrypted:
                warnings.append("Encrypted PDF detected; attempting decrypt with empty password.")
                try:
                    reader.decrypt("")
                except Exception:
                    raise ValueError("Cannot read encrypted PDF without password.")
                    
            total_pages = len(reader.pages)
            if total_pages == 0:
                warnings.append("PDF contains 0 pages.")
                
            for idx, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                clean_text = page_text.strip()
                if clean_text:
                    sections.append(
                        ExtractedSection(
                            text=clean_text,
                            page_number=idx + 1,
                            section_title=f"Page {idx + 1}"
                        )
                    )
            if not sections:
                warnings.append("No text extracted from PDF. Document might be scanned or image-only.")
        except Exception as e:
            raise ValueError(f"Failed to extract PDF text: {str(e)}")
            
    elif normalized_type in ["docx", "doc"]:
        try:
            doc = docx.Document(file_path)
            current_section = "Introduction"
            buffer = []
            
            for p in doc.paragraphs:
                p_text = p.text.strip()
                if not p_text:
                    continue
                # Heading detection
                if p.style and p.style.name and p.style.name.startswith("Heading"):
                    if buffer:
                        sections.append(
                            ExtractedSection(
                                text="\n".join(buffer),
                                page_number=None,
                                section_title=current_section
                            )
                        )
                        buffer = []
                    current_section = p_text
                else:
                    buffer.append(p_text)
                    
            if buffer:
                sections.append(
                    ExtractedSection(
                        text="\n".join(buffer),
                        page_number=None,
                        section_title=current_section
                    )
                )
                
            # Also extract tables
            for table_idx, table in enumerate(doc.tables):
                table_rows = []
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells]
                    table_rows.append(" | ".join(row_cells))
                if table_rows:
                    sections.append(
                        ExtractedSection(
                            text="\n".join(table_rows),
                            page_number=None,
                            section_title=f"Table {table_idx + 1}"
                        )
                    )
        except Exception as e:
            raise ValueError(f"Failed to extract DOCX text: {str(e)}")
            
    elif normalized_type in ["txt", "md", "markdown", "csv"]:
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
            clean_content = content.strip()
            if clean_content:
                sections.append(
                    ExtractedSection(
                        text=clean_content,
                        page_number=1,
                        section_title="Full Document"
                    )
                )
        except Exception as e:
            raise ValueError(f"Failed to extract plain text: {str(e)}")
    else:
        raise ValueError(f"Unsupported file format: {file_type}. Supported: PDF, DOCX, TXT")
        
    return sections, warnings
